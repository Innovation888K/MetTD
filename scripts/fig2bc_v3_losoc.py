#!/usr/bin/env python3
"""Fig2b/c on benchmark_v3 under the agreed protocol.

MetTD : every entry is scored independently (no training, no split). Statistics are
        taken over all 100 entries; the binary call is the manuscript convention
        "enzyme falls in the top 5% of the 1126-enzyme network".
ML    : XGBoost / SVM / BP / RF trained with leave-one-study-out CV (held-out study
        forms the test fold), out-of-fold probabilities pooled for ROC/AUC.
"""
import sys, os, json, warnings
sys.stdout.reconfigure(encoding='utf-8'); warnings.filterwarnings('ignore')
import numpy as np, pandas as pd

RD = 'D:/MetTD_recompute'
N = 9492
TOP_FRAC = 0.05
dist = np.memmap(f'{RD}/model_dist.bin', dtype='<f8', mode='r', shape=(N, N), order='F')
enz = pd.read_csv(f'{RD}/enzyme.csv', dtype=str)
enz_idx = enz['idx'].astype(int).to_numpy() - 1
enz_ec = enz['enzyme'].to_numpy()
IDH1 = int(np.where(enz_ec == '1.1.1.42')[0][0])
nE = len(enz_idx)

nm = pd.read_csv(f'{RD}/nodeid_map.csv', dtype=str).fillna('')
h2n = {}
for _, r in nm.iterrows():
    h = (r['HMDB.ID'] or '').strip().upper()
    if h.startswith('HMDB'):
        h2n.setdefault(h, set()).add(int(float(r['node_idx'])) - 1)
node2h = {}
for h, ns in h2n.items():
    for nd in ns:
        node2h.setdefault(nd, h)

tot = pd.read_excel(r'D:/wyc/代谢组数据库（metabolights）/abs数据库/totol_dataset.xlsx')
tot['HMDB'] = tot['HMDB.ID'].astype(str).str.strip().str.upper()
hm = tot['HMDB'].to_numpy()
nM = len(tot)
cols = {str(j): pd.to_numeric(tot[str(j)], errors='coerce').to_numpy() for j in range(1, 112)}

from openpyxl import load_workbook
wb = load_workbook(r'D:/wyc/代谢组数据库（metabolights）/database_detail.xlsx', read_only=True)
ws = wb[wb.sheetnames[0]]
dd = {}
for r in list(ws.iter_rows(values_only=True))[1:]:
    if r and r[0] is not None:
        dd[(str(r[1]), str(r[2]), str(r[3]))] = int(r[0])
wb.close()

v3 = pd.read_csv(f'{RD}/benchmark_v3.csv', encoding='utf-8-sig', dtype=str).fillna('')

def score_entry(nodes, levels):
    L = np.asarray(levels, float)
    s = L.std(ddof=1)
    if s == 0 or len(nodes) < 2:
        return None
    z2 = ((L - L.mean()) / s) ** 2
    R = np.asarray(nodes, int)
    D = dist[np.ix_(R, enz_idx)]
    W = np.where((D > 0) & np.isfinite(D), 1.0 / np.square(np.where(D == 0, 1, D)), 0.0)
    return z2 @ W

def f1_ref(prog):
    for c in (f'{RD}/ref_{prog}.csv', f'{RD}/ref_MTBLS3129_metformin.csv'):
        if os.path.exists(c):
            return c
    return None

rows = []
MET = np.zeros((len(v3), nM))
y = np.zeros(len(v3), int)
grp = []

for i, e in v3.iterrows():
    prog, st = e['Program'], e['status']
    y[i] = 1 if e['Target'] == 'IDH1' else 0
    grp.append(prog)
    if st == 'NEW-added':
        fp = f1_ref(prog)
        nodes = None
        if fp:
            d = pd.read_csv(fp)
            nodes = d['idx'].astype(int).to_numpy() - 1
            levels = d['level_meandiff'].to_numpy(float)
            for nd, lv in zip(nodes, levels):
                h = node2h.get(int(nd))
                if h:
                    j = np.where(hm == h)[0]
                    if len(j):
                        MET[i, int(j[0])] = lv
        if nodes is None:
            rows.append((prog, st, np.nan, np.nan, 'no-ref'))
            continue
        sc = score_entry(nodes, levels)
    else:
        oid = dd.get((prog, e['Ctrl_group'], e['Model_group']))
        col = cols.get(str(oid)) if oid is not None else None
        nodes = []
        levels = []
        if col is not None:
            for k, v in enumerate(col):
                if pd.notna(v) and abs(v) > 1e-6:
                    for nd in h2n.get(hm[k], ()):
                        nodes.append(nd)
                        levels.append(v)
            MET[i] = np.nan_to_num(col, nan=0.0)
        if len(nodes) < 2:
            rows.append((prog, st, np.nan, np.nan, 'no-data'))
            continue
        sc = score_entry(nodes, levels)
    if sc is None:
        rows.append((prog, st, np.nan, np.nan, 'degenerate'))
        continue
    rk = int((sc > sc[IDH1]).sum()) + 1
    rows.append((prog, st, rk, float(rk) / nE, ''))

E = pd.DataFrame(rows, columns=['Program', 'status', 'IDH1_rank', 'IDH1_pct', 'note'])
noval = E['IDH1_pct'].isna()
mettd = np.where(noval, 1.0, E['IDH1_pct'].to_numpy(float))
disc = -mettd
pred_m = (mettd <= TOP_FRAC).astype(int)
print(f"MetTD: scored {(~noval).sum()}/{len(E)}; unscored {int(noval.sum())} "
      f"{E.loc[noval,'note'].value_counts().to_dict()}")
print("IDH1 rank median: all "
      f"{E.loc[~noval,'IDH1_rank'].median():.0f}, pos "
      f"{E.loc[~noval & (y==1),'IDH1_rank'].median():.0f}, neg "
      f"{E.loc[~noval & (y==0),'IDH1_rank'].median():.0f} of {nE} enzymes")

from sklearn.preprocessing import StandardScaler
from sklearn.model_selection import LeaveOneGroupOut
from sklearn.metrics import roc_auc_score, roc_curve, accuracy_score, precision_score, \
    recall_score, f1_score, confusion_matrix
from sklearn.svm import SVC
from sklearn.ensemble import RandomForestClassifier
from sklearn.neural_network import MLPClassifier
from xgboost import XGBClassifier

grp = np.array(grp)
logo = list(LeaveOneGroupOut().split(MET, y, grp))
print(f"LOSO folds: {len(logo)} studies; train size "
      f"{min(len(tr) for tr, _ in logo)}-{max(len(tr) for tr, _ in logo)}")

def oof(make):
    p = np.full(len(y), np.nan)
    for tr, te in logo:
        sc = StandardScaler().fit(MET[tr])
        a = sc.transform(MET[tr])
        b = sc.transform(MET[te])
        m = make()
        m.fit(a, y[tr])
        p[te] = m.predict_proba(b)[:, 1]
    return p

P = {
 'MetTD': disc,
 'XGBoost': oof(lambda: XGBClassifier(n_estimators=35, learning_rate=0.1, max_depth=6,
                                      subsample=0.8, colsample_bytree=0.8,
                                      eval_metric='logloss', random_state=123)),
 'SVM': oof(lambda: SVC(kernel='sigmoid', gamma=5, probability=True, random_state=1)),
 'BP': oof(lambda: MLPClassifier(hidden_layer_sizes=(64, 32), max_iter=1000, random_state=1)),
 'RF': oof(lambda: RandomForestClassifier(n_estimators=100, max_features='sqrt', random_state=123)),
}

def metrics(yy, lab):
    tn, fp, fn, tp = confusion_matrix(yy, lab, labels=[0, 1]).ravel()
    return dict(acc=float(accuracy_score(yy, lab)),
                prec=float(precision_score(yy, lab, zero_division=0)),
                rec=float(recall_score(yy, lab)),
                f1=float(f1_score(yy, lab)),
                spec=float(tn / (tn + fp) if tn + fp else 0.0))

out = {'roc': {}, 'metrics': {}, 'protocol': {
    'mettd': 'per-entry F1 score, no split; call = IDH1 in top 5% of 1126 enzymes',
    'ml': f'leave-one-study-out CV, {len(logo)} folds, out-of-fold probabilities',
    'n_entries': int(len(y)), 'n_pos_IDH1': int(y.sum()), 'n_studies': int(len(logo))}}

for k, p in P.items():
    fpr, tpr, _ = roc_curve(y, p)
    auc = roc_auc_score(y, p)
    lab = pred_m if k == 'MetTD' else (p >= 0.5).astype(int)
    mm = metrics(y, lab)
    out['roc'][k] = {'fpr': fpr.tolist(), 'tpr': tpr.tolist()}
    out['metrics'][k] = {'auc': float(auc), **mm}
    print(f"{k:8s} AUC={auc:.3f} Acc={mm['acc']:.2f} Prec={mm['prec']:.2f} "
          f"Rec={mm['rec']:.2f} F1={mm['f1']:.2f} Spec={mm['spec']:.2f}")

E['y'] = y
E.to_csv(f'{RD}/fig2bc_v3_mettd_idh1rank.csv', index=False, encoding='utf-8-sig')
json.dump(out, open(f'{RD}/fig2bc_v3.json', 'w'))
print('wrote fig2bc_v3.json')
